"""Render structured discharge text as downloadable document formats."""

from __future__ import annotations

from dataclasses import dataclass
from html import escape
from io import BytesIO

from docx import Document
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import Flowable, Paragraph, SimpleDocTemplate, Spacer

from backend.formatting.structuring import SECTION_ORDER

DOCX_MEDIA_TYPE = (
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
)
MEDIA_TYPES = {
    "txt": "text/plain; charset=utf-8",
    "docx": DOCX_MEDIA_TYPE,
    "pdf": "application/pdf",
}


@dataclass(frozen=True)
class ExportArtifact:
    """Generated file content and the metadata needed for an HTTP download."""

    content: bytes
    media_type: str
    extension: str


def _render_docx(content: str) -> bytes:
    document = Document()
    for line in content.splitlines():
        if line in SECTION_ORDER:
            document.add_heading(line, level=1)
        elif line and set(line) != {"-"}:
            if line.startswith("- "):
                document.add_paragraph(line[2:], style="List Bullet")
            else:
                document.add_paragraph(line)
    output = BytesIO()
    document.save(output)
    return output.getvalue()


def _render_pdf(content: str) -> bytes:
    output = BytesIO()
    styles = getSampleStyleSheet()
    story: list[Flowable] = []
    for line in content.splitlines():
        if line in SECTION_ORDER:
            story.append(Paragraph(escape(line), styles["Heading2"]))
        elif line and set(line) != {"-"}:
            paragraph = line.removeprefix("- ")
            story.append(Paragraph(escape(paragraph), styles["BodyText"]))
        else:
            story.append(Spacer(1, 6))
    SimpleDocTemplate(output, pagesize=A4).build(story)
    return output.getvalue()


def create_export(content: str, output_format: str) -> ExportArtifact:
    """Render verified structured content without changing its wording."""
    if output_format == "txt":
        rendered = content.encode("utf-8")
    elif output_format == "docx":
        rendered = _render_docx(content)
    elif output_format == "pdf":
        rendered = _render_pdf(content)
    else:
        raise ValueError("Unsupported export format. Use PDF, DOCX, or TXT.")
    return ExportArtifact(
        content=rendered,
        media_type=MEDIA_TYPES[output_format],
        extension=output_format,
    )