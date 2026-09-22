"""Doctor-facing Streamlit interface for clinical discharge review."""
from __future__ import annotations

import io
import json
import os
import textwrap
import urllib.error
import urllib.request
from datetime import datetime
from html import escape
from typing import Any, Literal

import streamlit as st
from docx import Document
from pypdf import PdfReader

from frontend.samples import SAMPLE_DOCUMENTS

API = os.getenv("API_BASE_URL", "http://localhost:8000").rstrip("/")
Decision = Literal["accepted", "ignored"]
LABELS = {
    "diagnosis": "Diagnosis", "medications": "Medications",
    "follow_up_requirements": "Follow-up plan", "follow-up": "Follow-up plan",
    "warning_signs": "Warning signs", "warning signs": "Warning signs",
}

CSS = """
<style>
@import url('https://fonts.googleapis.com/css2?family=Manrope:wght@400;500;600;700&family=Sora:wght@400;600;700;800&display=swap');
:root{--paper:#F7FAF9;--surface:#fff;--ink:#182321;--muted:#64706D;--teal:#0F766E;--deep:#155E75;--red:#A43B3B;--line:rgba(24,35,33,.12);--soft:rgba(15,118,110,.08)}
html,body,[class*="css"]{font-family:Manrope,sans-serif}.stApp{background:var(--paper);color:var(--ink)}
h1,h2,h3,h4{font-family:Sora,sans-serif!important;letter-spacing:0!important}.block-container{max-width:none;padding:1rem 1.5rem 5rem}
[data-testid="stHeader"]{background:transparent}[data-testid="stSidebar"]{background:var(--paper);border-right:1px solid var(--line)}
[data-testid="stSidebar"] .block-container{padding:1.25rem .9rem}.brand{padding:.4rem .2rem 1rem}.brand-name{font-family:Sora;font-size:.98rem;font-weight:800}.brand-dot{color:var(--teal)}
.brand-subtitle{margin-top:.4rem;color:var(--muted);font-size:.62rem;font-weight:700;letter-spacing:.16em;text-transform:uppercase}
.title{display:flex;align-items:center;gap:.6rem;flex-wrap:wrap}.title h1{font-size:1.16rem;margin:0}.pill{padding:.18rem .48rem;border-radius:5px;background:var(--teal);color:#fff;font-size:.61rem;font-weight:700;letter-spacing:.1em;text-transform:uppercase}
.workflow{display:flex;justify-content:flex-end;align-items:center;gap:.34rem;flex-wrap:wrap}.step{font-size:.68rem;color:var(--muted);font-weight:600}.step.active{padding:.2rem .45rem;border-radius:4px;background:var(--ink);color:var(--paper)}.arrow{color:var(--teal)}
.score{display:flex;align-items:center;gap:.5rem;margin-top:.48rem}.track{width:9rem;height:.35rem;border-radius:99px;background:rgba(24,35,33,.1);overflow:hidden}.fill{height:100%;background:var(--teal)}.score span{font-size:.69rem;font-weight:700}.score small{font-size:.67rem;color:var(--muted)}
.panel{display:flex;justify-content:space-between;align-items:center;padding:.62rem 0;border-bottom:1px solid var(--line);margin-bottom:.3rem}.panel strong{font-family:Sora;font-size:.68rem;letter-spacing:.14em;text-transform:uppercase}.panel.final strong{color:var(--deep)}.panel span{font-size:.68rem;color:var(--muted)}
.section{min-height:8.3rem;padding:1rem .15rem 1.1rem;border-bottom:1px solid var(--line)}.section-head{display:flex;justify-content:space-between;margin-bottom:.5rem}.section-head h3{font-size:.87rem;margin:0}.section-head span{font-size:.61rem;font-weight:700;color:var(--muted);text-transform:uppercase}.copy{font-size:.81rem;line-height:1.65;color:rgba(24,35,33,.82);margin:0}.empty{color:var(--muted);font-style:italic}.approved{display:block;margin-top:.5rem;padding:.45rem .55rem;border-left:2px solid var(--teal);background:var(--soft);color:var(--teal)}
.suggestion{margin:.72rem 0 .2rem;padding:.78rem;border:1px solid rgba(21,94,117,.24);border-radius:8px;background:var(--surface)}.suggestion.accepted{border-color:rgba(15,118,110,.25);background:var(--soft)}.suggestion.ignored{border-color:rgba(164,59,59,.2);background:rgba(164,59,59,.06)}.suggestion-top{display:flex;justify-content:space-between}.state{font-size:.61rem;font-weight:800;letter-spacing:.12em;text-transform:uppercase;color:var(--deep)}.accepted .state{color:var(--teal)}.ignored .state{color:var(--red)}.suggestion p{font-size:.78rem;line-height:1.5;margin:.4rem 0 0}.suggestion .why{font-size:.69rem;color:var(--muted)}
.history{border-left:2px solid transparent;padding:.5rem .6rem;margin:.2rem 0;border-radius:6px}.history.active{background:var(--soft);border-left-color:var(--teal)}.history b{font-size:.72rem}.history div{font-size:.65rem;color:var(--muted);margin-top:.12rem}.reference{font-size:.67rem;line-height:1.45;color:var(--muted);padding-top:.7rem;border-top:1px solid var(--line)}
.empty-page{text-align:center;max-width:43rem;margin:7vh auto 0}.empty-page h1{font-size:1.6rem;margin-bottom:.45rem}.empty-page p{font-size:.84rem;line-height:1.6;color:var(--muted)}.footer{position:fixed;z-index:99;bottom:0;left:0;right:0;padding:.62rem 1.5rem;background:rgba(247,250,249,.96);border-top:1px solid var(--line);backdrop-filter:blur(10px);font-size:.7rem;color:var(--muted)}
.stButton>button,.stDownloadButton>button{border-radius:7px;min-height:2.25rem;font-family:Manrope;font-weight:700;font-size:.75rem;border-color:var(--line)}.stButton>button[kind="primary"],.stDownloadButton>button[kind="primary"]{background:var(--ink);color:var(--paper);border-color:var(--ink)}.stButton>button[kind="primary"]:hover,.stDownloadButton>button[kind="primary"]:hover{background:var(--teal);border-color:var(--teal)}
[data-testid="stFileUploaderDropzone"]{background:var(--surface);border:1px dashed rgba(15,118,110,.4);border-radius:8px}[data-testid="stExpander"]{border:0;border-top:1px solid var(--line);border-radius:0}
@media(max-width:900px){.block-container{padding:.7rem .8rem 6rem}.workflow{justify-content:flex-start;margin-top:.7rem}[data-testid="stHorizontalBlock"]{flex-wrap:wrap}[data-testid="column"]{min-width:100%!important;width:100%!important;flex:1 1 100%!important}.section{min-height:auto}.track{width:6rem}}
</style>
"""


def init() -> None:
    defaults = {"review": None, "source": "", "filename": "", "decisions": {}, "sessions": [], "export": False}
    for key, value in defaults.items():
        if key not in st.session_state:
            st.session_state[key] = value


def reset() -> None:
    for key, value in {"review": None, "source": "", "filename": "", "decisions": {}, "export": False}.items():
        st.session_state[key] = value


def request_json(path: str, method: str = "GET", body: bytes | None = None, headers: dict[str, str] | None = None) -> dict[str, Any]:
    request = urllib.request.Request(API + path, data=body, method=method, headers=headers or {})
    try:
        with urllib.request.urlopen(request, timeout=180) as response:
            return json.loads(response.read().decode())
    except urllib.error.HTTPError as error:
        try:
            detail = json.loads(error.read().decode()).get("detail", "The review request was rejected.")
        except (json.JSONDecodeError, UnicodeDecodeError):
            detail = "The review request was rejected."
        raise RuntimeError(str(detail)) from error
    except urllib.error.URLError as error:
        raise RuntimeError("The review service is unavailable. Confirm that FastAPI is running.") from error


def multipart(name: str, mime: str, data: bytes) -> tuple[bytes, str]:
    boundary = "----LumenReviewBoundary"
    prefix = f'--{boundary}\r\nContent-Disposition: form-data; name="file"; filename="{name.replace(chr(34), "")}"\r\nContent-Type: {mime}\r\n\r\n'.encode()
    return prefix + data + f"\r\n--{boundary}--\r\n".encode(), f"multipart/form-data; boundary={boundary}"


def source_text(name: str, data: bytes) -> str:
    suffix = name.lower().rsplit(".", 1)[-1]
    if suffix == "txt": return data.decode("utf-8", errors="replace")
    if suffix == "pdf": return "\n\n".join(page.extract_text() or "" for page in PdfReader(io.BytesIO(data)).pages).strip()
    if suffix == "docx": return "\n".join(paragraph.text for paragraph in Document(io.BytesIO(data)).paragraphs).strip()
    raise ValueError("Upload a PDF, DOCX, or TXT document.")


def analyze(name: str, mime: str, data: bytes) -> None:
    if len(data) > 5 * 1024 * 1024: raise ValueError("The file exceeds the 5 MB upload limit.")
    text = source_text(name, data)
    if not text.strip(): raise ValueError("No readable text was found in this document.")
    body, content_type = multipart(name, mime, data)
    review = request_json("/discharge/review", "POST", body, {"Content-Type": content_type})
    decisions = {int(item["suggestion_id"]): item["decision"] for item in review.get("suggestions", []) if isinstance(item.get("suggestion_id"), int) and item.get("decision")}
    st.session_state.review, st.session_state.source, st.session_state.filename = review, text, name
    st.session_state.decisions = decisions
    session = {"report_id": review.get("report_id"), "filename": name, "diagnosis": review.get("diagnosis", "Review"), "created": datetime.now().strftime("%d %b · %H:%M"), "review": review, "source": text, "decisions": decisions}
    st.session_state.sessions = [session, *[item for item in st.session_state.sessions if item.get("report_id") != session.get("report_id")]][:20]


def section_key(value: Any) -> str:
    return str(value or "").lower().strip().replace(" ", "_")


def sections(review: dict[str, Any]) -> list[dict[str, Any]]:
    output = []
    for index, item in enumerate(review.get("comparisons", []), 1):
        key = section_key(item.get("section", f"section-{index}"))
        output.append({"key": key, "label": LABELS.get(key, key.replace("_", " ").title()), "values": [str(value) for value in item.get("extracted_values", []) if value]})
    return output or [{"key": "diagnosis", "label": "Diagnosis", "values": [str(review.get("diagnosis", "Not identified"))]}]


def decide(suggestion_id: int, decision: Decision) -> None:
    report_id = st.session_state.review.get("report_id")
    if not isinstance(report_id, int): raise RuntimeError("This review has no saved report identifier.")
    request_json(f"/discharge/reports/{report_id}/suggestions/{suggestion_id}/decision", "POST", json.dumps({"decision": decision}).encode(), {"Content-Type": "application/json"})
    st.session_state.decisions[suggestion_id] = decision
    for item in st.session_state.sessions:
        if item.get("report_id") == report_id: item["decisions"] = dict(st.session_state.decisions)


def final_text(review: dict[str, Any]) -> str:
    result = ["FINAL REVIEWED DISCHARGE SUMMARY", ""]
    suggestions = review.get("suggestions", [])
    for section in sections(review):
        result += [section["label"].upper(), *(section["values"] or ["Not documented in the uploaded summary."])]
        for item in suggestions:
            sid = item.get("suggestion_id")
            if isinstance(sid, int) and st.session_state.decisions.get(sid) == "accepted" and section_key(item.get("section")) == section["key"]:
                result.append("Doctor-approved addition: " + str(item.get("explanation", "")))
        result.append("")
    return "\n".join(result).strip() + "\n"


def docx_data(text: str) -> bytes:
    document = Document()
    for index, block in enumerate(text.split("\n\n")):
        if index == 0 or (block.isupper() and "\n" not in block): document.add_heading(block, level=1 if index == 0 else 2)
        else: document.add_paragraph(block)
    output = io.BytesIO(); document.save(output); return output.getvalue()


def pdf_data(text: str) -> bytes:
    lines = []
    for line in text.splitlines(): lines.extend(textwrap.wrap(line, 88) or [""])
    pages = [lines[i:i + 48] for i in range(0, len(lines), 48)] or [[]]
    objects: list[bytes] = [b"", b"", b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>"]
    page_ids = []
    for page in pages:
        page_id, content_id = len(objects) + 1, len(objects) + 2; page_ids.append(page_id)
        commands = ["BT", "/F1 10 Tf", "54 756 Td", "13 TL"]
        for line in page:
            safe = line.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")
            commands += [f"({safe}) Tj", "T*"]
        stream = ("\n".join(commands + ["ET"])).encode("latin-1", errors="replace")
        objects += [f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Resources << /Font << /F1 3 0 R >> >> /Contents {content_id} 0 R >>".encode(), f"<< /Length {len(stream)} >>\nstream\n".encode() + stream + b"\nendstream"]
    objects[0] = b"<< /Type /Catalog /Pages 2 0 R >>"
    objects[1] = f"<< /Type /Pages /Kids [{' '.join(f'{item} 0 R' for item in page_ids)}] /Count {len(page_ids)} >>".encode()
    result = bytearray(b"%PDF-1.4\n%\xe2\xe3\xcf\xd3\n"); offsets = []
    for oid, value in enumerate(objects, 1): offsets.append(len(result)); result.extend(f"{oid} 0 obj\n".encode() + value + b"\nendobj\n")
    xref = len(result); result.extend(f"xref\n0 {len(objects)+1}\n0000000000 65535 f \n".encode())
    for offset in offsets: result.extend(f"{offset:010d} 00000 n \n".encode())
    result.extend(f"trailer\n<< /Size {len(objects)+1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF".encode()); return bytes(result)


def sidebar() -> None:
    with st.sidebar:
        st.markdown('<div class="brand"><div class="brand-name">Lumen<span class="brand-dot">.</span>Review</div><div class="brand-subtitle">Discharge workspace</div></div>', unsafe_allow_html=True)
        if st.button("New review", type="primary", use_container_width=True): reset(); st.rerun()
        query = st.text_input("Search", placeholder="Search history", label_visibility="collapsed").casefold().strip()
        st.selectbox("Filter", ["All reviews", "In review", "Complete"], label_visibility="collapsed")
        st.caption("RECENT REVIEWS")
        visible = [item for item in st.session_state.sessions if not query or query in (str(item["filename"]) + str(item["diagnosis"])).casefold()]
        if not visible: st.caption("No matching reviews in this session.")
        for index, item in enumerate(visible):
            count = len(item["review"].get("suggestions", [])); open_count = max(count - len(item["decisions"]), 0)
            active = item.get("report_id") == (st.session_state.review or {}).get("report_id")
            st.markdown(f'<div class="history {"active" if active else ""}"><b>{escape(item["filename"])}</b><div>{escape(str(item["diagnosis"]))} · {open_count} open · {item["created"]}</div></div>', unsafe_allow_html=True)
            if st.button("Open", key=f"open-{index}-{item.get('report_id')}", use_container_width=True):
                st.session_state.review, st.session_state.source, st.session_state.filename = item["review"], item["source"], item["filename"]
                st.session_state.decisions = dict(item["decisions"]); st.rerun()
        st.markdown('<div class="reference">Reference content from MedlinePlus. Patient health information for comparison, not treatment guidance.</div>', unsafe_allow_html=True)


def upload_view() -> None:
    st.markdown('<div class="empty-page"><h1>Review a discharge summary</h1><p>Upload a synthetic or de-identified document. Every suggestion stays pending until a doctor explicitly accepts or ignores it.</p></div>', unsafe_allow_html=True)
    _, center, _ = st.columns([1, 1.35, 1])
    with center:
        upload = st.file_uploader("Upload discharge summary", type=["pdf", "docx", "txt"], help="5 MB maximum. Do not upload identifiable patient information.")
        st.caption("PDF, DOCX, or TXT · 5 MB maximum · De-identified documents only")
        if upload and st.button("Analyze document", type="primary", use_container_width=True):
            try:
                with st.status("Analyzing the discharge summary…", expanded=True) as status:
                    st.write("Reading documented clinical fields"); st.write("Finding the closest available health reference"); st.write("Preparing suggestions for doctor review")
                    analyze(upload.name, upload.type or "application/octet-stream", upload.getvalue()); status.update(label="Review ready", state="complete", expanded=False)
                st.rerun()
            except (ValueError, RuntimeError) as error: st.error(str(error))
        st.divider(); st.caption("TRY A SYNTHETIC SAMPLE")
        for column, (condition, sample) in zip(st.columns(3), SAMPLE_DOCUMENTS.items()):
            with column:
                if st.button(condition, use_container_width=True):
                    try: analyze(f"sample-{condition.lower()}.txt", "text/plain", sample.encode()); st.rerun()
                    except (ValueError, RuntimeError) as error: st.error(str(error))


def workflow(review: dict[str, Any]) -> None:
    score = review.get("completeness_score") if isinstance(review.get("completeness_score"), int) else 0
    pending = sum(item.get("suggestion_id") not in st.session_state.decisions for item in review.get("suggestions", [])); current = "Review" if pending else "Approve"
    title, flow = st.columns([1.05, 1])
    with title: st.markdown(f'<div class="title"><h1>{escape(st.session_state.filename)}</h1><span class="pill">{"In review" if pending else "Ready to finalize"}</span></div><div class="score"><div class="track"><div class="fill" style="width:{score}%"></div></div><span>Completeness {score}%</span><small>{escape(str(review.get("diagnosis", "")))} reference match</small></div>', unsafe_allow_html=True)
    with flow:
        steps = ["Upload", "Analyze", "Review", "Approve", "Final", "Export"]
        markup = '<div class="workflow">' + ''.join(f'<span class="step{" active" if step == current else ""}">{step}</span>{"<span class=arrow>→</span>" if index < 5 else ""}' for index, step in enumerate(steps)) + '</div>'
        st.markdown(markup, unsafe_allow_html=True)
    st.caption("Patient health reference: MedlinePlus · Educational comparison only, not treatment guidance"); st.divider()


def original_section(section: dict[str, Any], number: int) -> None:
    copy = "<br>".join(escape(value) for value in section["values"]) or '<span class="empty">Not documented in the uploaded summary.</span>'
    st.markdown(f'<div class="section"><div class="section-head"><h3>{escape(section["label"])}</h3><span>Section {number}</span></div><p class="copy">{copy}</p></div>', unsafe_allow_html=True)


def final_section(section: dict[str, Any], number: int, items: list[dict[str, Any]]) -> None:
    copy = "<br>".join(escape(value) for value in section["values"]) or '<span class="empty">Not documented in the uploaded summary.</span>'
    for item in items:
        sid = item.get("suggestion_id")
        if isinstance(sid, int) and st.session_state.decisions.get(sid) == "accepted": copy += f'<span class="approved">{escape(str(item.get("explanation", "")))}</span>'
    st.markdown(f'<div class="section"><div class="section-head"><h3>{escape(section["label"])}</h3><span>Section {number}</span></div><p class="copy">{copy}</p>', unsafe_allow_html=True)
    for item in items:
        sid = item.get("suggestion_id")
        if not isinstance(sid, int): continue
        state = st.session_state.decisions.get(sid, "pending"); label = {"pending":"Suggestion · pending","accepted":"Accepted","ignored":"Ignored"}[state]
        st.markdown(f'<div class="suggestion {state}"><div class="suggestion-top"><span class="state">{label}</span><span class="history-meta">Evidence-backed</span></div><p>{escape(str(item.get("explanation", "")))}</p><p class="why"><strong>Why:</strong> {escape(str(item.get("guideline_passage", "")))}</p></div>', unsafe_allow_html=True)
        if state == "pending":
            accept_col, ignore_col = st.columns(2)
            with accept_col:
                if st.button("Accept", key=f"accept-{sid}", type="primary", use_container_width=True):
                    try: decide(sid, "accepted"); st.rerun()
                    except RuntimeError as error: st.error(str(error))
            with ignore_col:
                if st.button("Ignore", key=f"ignore-{sid}", use_container_width=True):
                    try: decide(sid, "ignored"); st.rerun()
                    except RuntimeError as error: st.error(str(error))
        else: st.caption(f"Decision recorded as {state}. The audit log retains this decision.")
    st.markdown("</div>", unsafe_allow_html=True)


def export_view(review: dict[str, Any]) -> None:
    text = final_text(review); stem = st.session_state.filename.rsplit(".", 1)[0] or "discharge-summary"
    st.markdown("#### Export final reviewed document"); st.caption("Only doctor-accepted additions are included. Ignored and pending suggestions remain excluded.")
    pdf_col, docx_col, txt_col = st.columns(3)
    with pdf_col: st.download_button("PDF", pdf_data(text), stem + "-final.pdf", "application/pdf", type="primary", use_container_width=True)
    with docx_col: st.download_button("DOCX", docx_data(text), stem + "-final.docx", "application/vnd.openxmlformats-officedocument.wordprocessingml.document", use_container_width=True)
    with txt_col: st.download_button("TXT", text.encode(), stem + "-final.txt", "text/plain", use_container_width=True)


def review_view() -> None:
    review = st.session_state.review
    if not review: upload_view(); return
    workflow(review); compared = sections(review); suggestions = review.get("suggestions", []); known = {item["key"] for item in compared}
    left, right = st.columns(2, gap="medium")
    with left:
        st.markdown('<div class="panel"><strong>Original</strong><span>Uploaded document</span></div>', unsafe_allow_html=True)
        for number, section in enumerate(compared, 1): original_section(section, number)
        with st.expander("View extracted source text"): st.text(st.session_state.source)
    with right:
        label_col, button_col = st.columns([2, 1])
        with label_col: st.markdown('<div class="panel final"><strong>Final · accepted changes</strong><span>Live preview</span></div>', unsafe_allow_html=True)
        with button_col:
            if st.button("Export", type="primary", use_container_width=True): st.session_state.export = not st.session_state.export
        for number, section in enumerate(compared, 1): final_section(section, number, [item for item in suggestions if section_key(item.get("section")) == section["key"]])
        unmatched = [item for item in suggestions if section_key(item.get("section")) not in known]
        if unmatched: final_section({"key":"additional","label":"Additional findings","values":[]}, len(compared)+1, unmatched)
        if st.session_state.export: st.divider(); export_view(review)
    counts = {"pending":0,"accepted":0,"ignored":0}
    for item in suggestions: counts[st.session_state.decisions.get(item.get("suggestion_id"), "pending")] += 1
    st.markdown(f'<div class="footer"><strong>Review progress:</strong> {counts["pending"]} pending · {counts["accepted"]} accepted · {counts["ignored"]} ignored</div>', unsafe_allow_html=True)


def main() -> None:
    st.set_page_config(page_title="Lumen Review · Clinical Discharge Review", page_icon="L", layout="wide", initial_sidebar_state="expanded")
    init(); st.markdown(CSS, unsafe_allow_html=True); sidebar(); review_view()


if __name__ == "__main__": main()
